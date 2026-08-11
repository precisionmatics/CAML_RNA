"""
Convert manuscript_caml_rna.tex to a Word DOCX matching the style of the
existing revision DOCX (Revisions/07_July_2026_Revised_CAML_RNA_manuscript.docx).
"""

import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

TEX_FILE  = "/home/stalin/Desktop/CAML/manuscript_caml_rna.tex"
OUT_FILE  = "/home/stalin/Desktop/CAML/manuscript_caml_rna_revised.docx"

# ── LaTeX → plain text helpers ────────────────────────────────────────────────

def clean_latex(text):
    """Strip LaTeX commands and convert to readable plain text."""
    # Comments
    text = re.sub(r'%.*', '', text)
    # Non-breaking space
    text = text.replace('~', ' ')
    # Dashes
    text = text.replace('---', '—').replace('--', '–')
    # Quotes
    text = text.replace("``", '“').replace("''", '”')
    text = text.replace("`", '‘').replace("'", '’')
    # Common formatting
    text = re.sub(r'\\textbf\{([^}]*)\}', r'\1', text)
    text = re.sub(r'\\textit\{([^}]*)\}', r'\1', text)
    text = re.sub(r'\\emph\{([^}]*)\}', r'\1', text)
    text = re.sub(r'\\text\{([^}]*)\}', r'\1', text)
    text = re.sub(r'\\mathrm\{([^}]*)\}', r'\1', text)
    text = re.sub(r'\\mathbf\{([^}]*)\}', r'\1', text)
    text = re.sub(r'\\bm\{([^}]*)\}', r'\1', text)
    text = re.sub(r'\\mathcal\{([^}]*)\}', r'\1', text)
    # Citations
    text = re.sub(r'\\cite\{([^}]*)\}', lambda m: '[' + m.group(1).split(',')[0].strip() + ']', text)
    # Refs
    text = re.sub(r'\\ref\{[^}]*\}', '[X]', text)
    text = re.sub(r'\\eqref\{[^}]*\}', '(X)', text)
    text = re.sub(r'\\label\{[^}]*\}', '', text)
    # Special commands
    text = re.sub(r'\\url\{([^}]*)\}', r'\1', text)
    text = re.sub(r'\\kD\b', 'pKd', text)
    text = re.sub(r'\\ES\b', 'ES', text)
    text = re.sub(r'\\CS\b', 'CS', text)
    text = re.sub(r'\\PSRT\b', 'PSRT', text)
    text = re.sub(r'\\VR\b', 'VR', text)
    text = re.sub(r'\\Tor\b', 'Tor', text)
    text = re.sub(r'\\dR\b', 'Δr', text)
    # Math mode — preserve inline math readable
    def clean_math(m):
        s = m.group(1)
        s = re.sub(r'\\([a-zA-Z]+)', lambda x: x.group(1), s)
        s = s.replace('{', '').replace('}', '')
        s = s.replace('_', '').replace('^', '')
        s = s.strip()
        return s
    text = re.sub(r'\$\$(.+?)\$\$', lambda m: '[' + clean_math(m) + ']', text, flags=re.DOTALL)
    text = re.sub(r'\$(.+?)\$', clean_math, text)
    # Accents and special chars
    text = re.sub(r'\\\'([a-zA-Z])', r'\1', text)
    text = re.sub(r'\\`([a-zA-Z])', r'\1', text)
    text = re.sub(r'\\"([a-zA-Z])', r'\1', text)
    text = re.sub(r'\\,', ' ', text)
    text = re.sub(r'\\;', ' ', text)
    text = re.sub(r'\\:', ' ', text)
    text = re.sub(r'\\!', '', text)
    text = re.sub(r'\\ ', ' ', text)
    text = re.sub(r'\\quad\b', '  ', text)
    text = re.sub(r'\\qquad\b', '   ', text)
    text = re.sub(r'\\noindent\b', '', text)
    text = re.sub(r'\\par\b', '', text)
    text = re.sub(r'\\item\b', '• ', text)
    text = re.sub(r'\\small\b', '', text)
    text = re.sub(r'\\footnotesize\b', '', text)
    text = re.sub(r'\\normalsize\b', '', text)
    text = re.sub(r'\\textbf\b', '', text)
    text = re.sub(r'\\textit\b', '', text)
    text = re.sub(r'\\vspace\*?\{[^}]*\}', '', text)
    text = re.sub(r'\\hspace\*?\{[^}]*\}', '', text)
    text = re.sub(r'\\newline\b', '\n', text)
    text = re.sub(r'\\\\', '\n', text)
    # Remove remaining backslash commands
    text = re.sub(r'\\[a-zA-Z@]+(\*)?(\{[^}]*\})?', '', text)
    # Clean up braces
    text = text.replace('{', '').replace('}', '')
    # Collapse whitespace
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


def eq_to_text(eq_body):
    """Render a display equation block as a readable string."""
    eq = eq_body.strip()
    eq = re.sub(r'\\label\{[^}]*\}', '', eq)
    eq = re.sub(r'\\nonumber', '', eq)
    # Clean up
    eq = re.sub(r'\\begin\{[^}]+\}', '', eq)
    eq = re.sub(r'\\end\{[^}]+\}', '', eq)
    eq = re.sub(r'\\\\', '  ', eq)
    eq = re.sub(r'&', ' ', eq)
    eq = clean_latex(eq)
    eq = re.sub(r'\s+', ' ', eq).strip()
    return eq


# ── Parse the .tex file into a list of blocks ─────────────────────────────────

class Block:
    def __init__(self, kind, content, level=0):
        self.kind = kind      # 'heading', 'para', 'equation', 'table', 'figure', 'list'
        self.content = content
        self.level = level    # heading level (1=section, 2=sub, 3=subsub, 4=para)


def parse_tex(tex):
    blocks = []

    # ── Special handling: hardcode title block from known manuscript ──
    blocks.append(Block('heading',
        'CAML-RNA: Commutative Algebra Machine Learning for RNA–Ligand Binding Affinity Prediction',
        0))
    blocks.append(Block('para',
        'Stalin Arulsamy,¹ Yashwanth Krishna,² Rajesh Kumar,³ and Vanktesh Kumar¹*'))
    blocks.append(Block('para',
        '¹Department of Pharmaceutical Chemistry, School of Pharmaceutical Sciences, '
        'Lovely Professional University, Phagwara, Punjab, India  '
        '²Department of Pharmacy Practice, Manipal College of Pharmaceutical Sciences, '
        'Manipal, Karnataka, India  '
        '³Department of Pharmacy Practice, School of Pharmaceutical Sciences, '
        'Lovely Professional University, Phagwara, Punjab, India  '
        '*Email: Vankteshkumar555@hotmail.com'))

    # Extract abstract from twocolumn block
    twocol_m = re.search(r'\\twocolumn\[(.*?)\]\s*%%\s*end', tex, re.DOTALL)
    if twocol_m:
        hdr = twocol_m.group(1)
        abs_m = re.search(r'\\textbf\{Abstract\.\}\\quad\s*(.*?)\\par\}', hdr, re.DOTALL)
        if abs_m:
            abstract = clean_latex(abs_m.group(1)).replace('\n', ' ').strip()
            blocks.append(Block('heading', 'Abstract', 1))
            blocks.append(Block('para', abstract))

    # Strip preamble: everything up to and including \begin{document}
    body_m = re.search(r'\\begin\{document\}(.*?)\\end\{document\}', tex, re.DOTALL)
    if body_m:
        tex_body = body_m.group(1)
        # Remove the twocolumn header block
        tex_body = re.sub(r'\\twocolumn\[.*?\]\s*%%\s*end.*?\n', '', tex_body, flags=re.DOTALL)
    else:
        tex_body = tex

    lines  = tex_body.split('\n')
    i = 0
    buf = []

    def flush_buf():
        text = ' '.join(buf).strip()
        buf.clear()
        if text:
            text = clean_latex(text)
            if text.strip():
                blocks.append(Block('para', text))

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Skip lines we don't want
        if not stripped or stripped.startswith('%'):
            flush_buf() if not stripped else None
            i += 1
            continue

        # Comments
        if stripped.startswith('%'):
            i += 1
            continue

        # Section headings
        m = re.match(r'\\(section|subsection|subsubsection|paragraph)\*?\{(.+)\}', stripped)
        if m:
            flush_buf()
            kind = m.group(1)
            lvl = {'section': 1, 'subsection': 2, 'subsubsection': 3, 'paragraph': 4}[kind]
            title = clean_latex(m.group(2))
            blocks.append(Block('heading', title, lvl))
            i += 1
            continue

        # Display equations
        if re.search(r'\\begin\{(equation|align|equation\*|align\*|multline)\*?\}', stripped):
            flush_buf()
            env_match = re.search(r'\\begin\{([^}]+)\}', stripped)
            env = env_match.group(1) if env_match else 'equation'
            eq_lines = [stripped]
            i += 1
            while i < len(lines) and not re.search(r'\\end\{' + re.escape(env.replace('*','')) + r'\*?\}', lines[i]):
                eq_lines.append(lines[i])
                i += 1
            eq_lines.append(lines[i] if i < len(lines) else '')
            eq_text = eq_to_text('\n'.join(eq_lines))
            if eq_text.strip():
                blocks.append(Block('equation', eq_text))
            i += 1
            continue

        # Tables
        if re.search(r'\\begin\{table', stripped):
            flush_buf()
            tbl_lines = [stripped]
            i += 1
            while i < len(lines) and '\\end{table' not in lines[i]:
                tbl_lines.append(lines[i])
                i += 1
            tbl_lines.append(lines[i] if i < len(lines) else '')
            blocks.append(Block('table', '\n'.join(tbl_lines)))
            i += 1
            continue

        # Figures
        if re.search(r'\\begin\{figure', stripped):
            flush_buf()
            fig_lines = []
            i += 1
            while i < len(lines) and '\\end{figure' not in lines[i]:
                fig_lines.append(lines[i])
                i += 1
            # Extract caption
            cap = ''
            cap_text = ' '.join(fig_lines)
            cm = re.search(r'\\caption\{(.+)', cap_text, re.DOTALL)
            if cm:
                cap = cm.group(1).split('\\label')[0].strip().rstrip('}')
            blocks.append(Block('figure', clean_latex(cap)))
            i += 1
            continue

        # itemize / enumerate
        if re.search(r'\\begin\{(itemize|enumerate)', stripped):
            flush_buf()
            i += 1
            while i < len(lines) and not re.search(r'\\end\{(itemize|enumerate)', lines[i]):
                item_line = lines[i].strip()
                if item_line.startswith('\\item'):
                    item_text = clean_latex(item_line[5:].strip())
                    if item_text:
                        blocks.append(Block('list', item_text))
                i += 1
            i += 1
            continue

        # Bibliography entries — include as reference list
        if stripped.startswith('\\bibitem'):
            flush_buf()
            ref_lines = [stripped]
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('\\bibitem') and \
                  not lines[i].strip().startswith('\\end{thebibliography}'):
                ref_lines.append(lines[i])
                i += 1
            ref_text = ' '.join(ref_lines)
            ref_text = clean_latex(ref_text)
            if ref_text.strip():
                blocks.append(Block('para', ref_text.strip()))
            continue

        # Skip begin/end of known environments we've handled
        if re.search(r'\\(begin|end)\{(thebibliography|tabular|document|figure|table|equation|align|itemize|enumerate|abstract)\*?\}', stripped):
            i += 1
            continue

        # Normal text lines — accumulate into buffer
        if stripped and not stripped.startswith('\\toprule') and \
           not stripped.startswith('\\midrule') and not stripped.startswith('\\bottomrule') and \
           not stripped.startswith('\\cmidrule') and not stripped.startswith('\\hline'):
            # Skip pure LaTeX command lines
            if re.match(r'^\\[a-zA-Z]+(\[.*\])?\{.*\}$', stripped) and \
               not any(c in stripped for c in ['caption', 'label', 'small', 'footnote']):
                i += 1
                continue
            buf.append(stripped)
        elif not stripped:
            flush_buf()

        i += 1

    flush_buf()
    return blocks


# ── Parse tabular environment → list of rows ─────────────────────────────────

def parse_tabular(tbl_text):
    """Extract caption and rows from a table block."""
    # Caption
    cap = ''
    cm = re.search(r'\\caption\{(.+?)(?:\}|\n)', tbl_text, re.DOTALL)
    if cm:
        cap = clean_latex(cm.group(1).split('\\label')[0])

    # Find tabular body
    tab_m = re.search(r'\\begin\{tabular\}(\{[^}]+\})?(.*?)\\end\{tabular\}', tbl_text, re.DOTALL)
    if not tab_m:
        return cap, []

    body = tab_m.group(2)
    # Remove hline-type commands and col formatting
    body = re.sub(r'\\(toprule|midrule|bottomrule|hline)[^\n]*', '', body)
    body = re.sub(r'\\cmidrule[^\n]*', '', body)
    body = re.sub(r'\\multirow\{[^}]+\}\{[^}]+\}\{([^}]*)\}', r'\1', body)
    body = re.sub(r'\\multicolumn\{[^}]+\}\{[^}]+\}\{([^}]*)\}', r'\1', body)

    rows = []
    for line in body.split('\\\\'):
        line = line.strip().rstrip('\\').strip()
        if not line:
            continue
        cells = [clean_latex(c.strip()) for c in line.split('&')]
        if any(c.strip() for c in cells):
            rows.append(cells)

    return cap, rows


# ── Build the Word document ───────────────────────────────────────────────────

def build_docx(blocks):
    doc = Document()

    # Page margins
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    section = doc.sections[0]
    section.top_margin    = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin   = Inches(1.25)
    section.right_margin  = Inches(1.25)

    # Default font
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)

    def add_heading(text, level):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        if level == 0:  # title
            run.font.size = Pt(16)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after  = Pt(8)
        elif level == 1:
            run.font.size = Pt(12)
            run.italic    = True
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after  = Pt(4)
        elif level == 3:
            run.font.size = Pt(12)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after  = Pt(2)
        else:
            run.font.size = Pt(12)
            p.paragraph_format.space_before = Pt(4)

    def add_para(text, italic=False, size=12, space_before=0):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.italic = italic
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after  = Pt(4)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    def add_equation(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(11)
        run.italic = True
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after  = Pt(4)

    in_biblio = False

    for block in blocks:
        if block.kind == 'heading':
            if 'References' in block.content or 'Bibliography' in block.content:
                in_biblio = True
            add_heading(block.content, block.level)

        elif block.kind == 'para':
            text = block.content.strip()
            if not text:
                continue
            sz = 10 if in_biblio else 12
            add_para(text, size=sz)

        elif block.kind == 'equation':
            if block.content.strip():
                add_equation(block.content)

        elif block.kind == 'list':
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(block.content)
            run.font.size = Pt(12)

        elif block.kind == 'figure':
            p = doc.add_paragraph()
            run = p.add_run('[FIGURE] ')
            run.bold = True
            run.font.size = Pt(11)
            run2 = p.add_run(block.content)
            run2.font.size = Pt(10)
            run2.italic = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after  = Pt(6)

        elif block.kind == 'table':
            cap, rows = parse_tabular(block.content)
            if cap:
                p = doc.add_paragraph()
                r = p.add_run('Table: ' + cap)
                r.bold = True
                r.font.size = Pt(11)
                p.paragraph_format.space_before = Pt(8)

            if rows:
                # Determine max columns
                n_cols = max(len(r) for r in rows)
                tbl = doc.add_table(rows=len(rows), cols=n_cols)
                tbl.style = 'Table Grid'
                for ri, row in enumerate(rows):
                    for ci, cell_text in enumerate(row):
                        if ci < n_cols:
                            cell = tbl.cell(ri, ci)
                            cell.text = cell_text
                            # Make header row bold
                            if ri == 0:
                                for run in cell.paragraphs[0].runs:
                                    run.bold = True
                doc.add_paragraph()  # space after table

    return doc


def main():
    print(f"Reading {TEX_FILE}...")
    with open(TEX_FILE, 'r', encoding='utf-8') as f:
        tex = f.read()

    print("Parsing LaTeX...")
    blocks = parse_tex(tex)
    print(f"  {len(blocks)} blocks extracted")

    print("Building DOCX...")
    doc = build_docx(blocks)

    doc.save(OUT_FILE)
    print(f"Saved: {OUT_FILE}")


if __name__ == '__main__':
    main()
