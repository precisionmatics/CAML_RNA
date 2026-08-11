#!/usr/bin/env python3
"""
fix_refs_and_renumber.py
Apply CrossRef-verified corrections, then renumber citations.

Verified fixes:
  [27] Song J et al. DeepMIF (correct authors/title/year, DOI 10.1021/acs.jcim.5c02946)
  [28] Huang Z et al. DeepRSMA (not "DeepMIF", DOI 10.1093/bioinformatics/btae678)
  [29] EMMPTNet — not found in CrossRef: hallucinated → REMOVE
  [32] RDKit — malformed DOI → fix format
  [38] Kramer O scikit-learn → Pedregosa et al. JMLR 2011

Table fixes:
  Row 6 EMMPTNet [29] → DeepMIF [27]  (Song J JCIM 2026, NA-L)
  Row 7 DeepMIF [27]  → DeepRSMA [28] (Huang Z Bioinformatics 2024, R-SIM)
"""

import re
from docx import Document

DOCX_IN  = "Revisions/10_August_2026_Revised_CAML_RNA_manuscript.docx"
DOCX_OUT = "Revisions/11_August_2026_Revised_CAML_RNA_manuscript.docx"

doc = Document(DOCX_IN)
paras = list(doc.paragraphs)
ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

# ── Find reference section ──────────────────────────────────────────────────
ref_start = next(i for i, p in enumerate(paras)
                 if re.match(r'^[Rr]eferences?\s*$', p.text.strip()))

# ── 1. Fix reference [27]: DeepMIF correct (Song J et al.) ──────────────────
REF27_NEW = ("[27]  Song J, Gao A, Tian S, Yang Q, Deng L, Feng Q, Hou M. "
             "DeepMIF: A Multiview Interactive Fusion-Based Deep Learning Method for RNA–"
             "Small Molecule Binding Affinity Prediction. "
             "Journal of Chemical Information and Modeling. 2026;66(7):3575-89.")

# ── 2. Fix reference [28]: DeepRSMA (not DeepMIF) ──────────────────────────
REF28_NEW = ("[28]  Huang Z, Wang Y, Chen S, Tan YS, Deng L, Wu M. "
             "DeepRSMA: a cross-fusion-based deep learning method for RNA–"
             "small molecule binding affinity prediction. "
             "Bioinformatics. 2024 Dec;40(12):btae678.")

# ── 3. Fix reference [32]: RDKit correct format ─────────────────────────────
REF32_NEW = ("[32]  Landrum G. RDKit: Open-Source Cheminformatics. "
             "2024. Available from: https://www.rdkit.org. "
             "DOI: 10.5281/zenodo.591637.")

# ── 4. Fix reference [38]: Pedregosa et al. (scikit-learn) ─────────────────
REF38_NEW = ("[38]  Pedregosa F, Varoquaux G, Gramfort A, Michel V, Thirion B, Grisel O, "
             "Blondel M, Prettenhofer P, Weiss R, Dubourg V, Vanderplas J, Passos A, "
             "Cournapeau D, Brucher M, Perrot M, Duchesnay E. Scikit-learn: Machine "
             "Learning in Python. Journal of Machine Learning Research. 2011;12(85):2825-30.")

def set_para_text(p, new_text):
    """Replace all runs in paragraph with new_text in the first run, clear rest."""
    if not p.runs:
        return
    p.runs[0].text = new_text
    for run in p.runs[1:]:
        run.text = ''

def find_ref_para(num):
    """Find the reference list paragraph that starts with [num]."""
    for p in paras[ref_start + 1:]:
        if p.text.strip().startswith(f'[{num}]'):
            return p
    return None

print("Fixing reference texts...")
p27 = find_ref_para(27)
p28 = find_ref_para(28)
p29 = find_ref_para(29)
p32 = find_ref_para(32)
p38 = find_ref_para(38)

if p27: set_para_text(p27, REF27_NEW); print("  [27] DeepMIF corrected (Song J et al.)")
if p28: set_para_text(p28, REF28_NEW); print("  [28] DeepRSMA title fixed (was incorrectly 'DeepMIF')")
if p32: set_para_text(p32, REF32_NEW); print("  [32] RDKit DOI fixed")
if p38: set_para_text(p38, REF38_NEW); print("  [38] Scikit-learn → Pedregosa et al. JMLR 2011")

# ── 5. Remove [29] from reference list ──────────────────────────────────────
if p29:
    body = doc.element.body
    body.remove(p29._element)
    print("  [29] EMMPTNet removed (not found in CrossRef)")
else:
    print("  WARNING: [29] not found")

# ── 6. Update Table 1 ────────────────────────────────────────────────────────
print("\nUpdating Table 1...")
tbl = doc.tables[0]

# Row 6: EMMPTNet [29] → DeepMIF [27]
row6_cell = tbl.rows[6].cells[0]
old6 = row6_cell.text
if 'EMMPTNet' in old6:
    for p in row6_cell.paragraphs:
        for run in p.runs:
            if 'EMMPTNet' in run.text:
                run.text = run.text.replace('EMMPTNet [29]', 'DeepMIF [27]')
            elif '[29]' in run.text:
                run.text = run.text.replace('[29]', '[27]')
    print(f"  Row 6: '{old6.strip()}' → '{row6_cell.text.strip()}'")

# Row 7: DeepMIF [27] → DeepRSMA [28]
row7_cell = tbl.rows[7].cells[0]
old7 = row7_cell.text
if 'DeepMIF' in old7:
    for p in row7_cell.paragraphs:
        for run in p.runs:
            if 'DeepMIF [27]' in run.text:
                run.text = run.text.replace('DeepMIF [27]', 'DeepRSMA [28]')
            elif 'DeepMIF' in run.text:
                run.text = run.text.replace('DeepMIF', 'DeepRSMA')
            elif '[27]' in run.text:
                run.text = run.text.replace('[27]', '[28]')
    print(f"  Row 7: '{old7.strip()}' → '{row7_cell.text.strip()}'")

# ── 7. Update body text ──────────────────────────────────────────────────────
print("\nUpdating body text...")

# Para 024: fix EMMPTNet and DeepMIF references
p24 = paras[24]
for run in p24.runs:
    txt = run.text
    # "DeepMIF [27,28] were trained on the R-SIM" → "DeepRSMA [28] were evaluated on the R-SIM"
    txt = txt.replace(
        'RSAPred [17] and DeepMIF [27,28] were trained on the R-SIM database; '
        'EMMPTNet [29] was evaluated on the NA-L database under 10-fold CV, '
        'which contains approximately 10 times more binding measurements than NA-L does, '
        'encompasses diverse assay types, and has a different RNA subtype distribution.',
        'RSAPred [17] and DeepRSMA [28] were evaluated on the R-SIM database, '
        'which contains approximately 10 times more RNA–ligand binding measurements '
        'than NA-L and covers a different subtype distribution; '
        'DeepMIF [27] was evaluated on the NA-L database under 10-fold CV.'
    )
    if txt != run.text:
        run.text = txt
        print("  Para 024: RSAPred/DeepMIF/EMMPTNet sentence updated")
        break

# Fallback: piecewise fixes if full-string replacement didn't match
p24_text = p24.text
if 'EMMPTNet' in p24_text:
    for run in p24.runs:
        run.text = (run.text
            .replace('EMMPTNet [29]', 'DeepMIF [27]')
            .replace('DeepMIF [27,28]', 'DeepRSMA [28]'))
    print("  Para 024 (fallback piecewise fix)")

# Para 027: "DeepMIF [27] (R = 0.784 on R-SIM)" → "DeepRSMA [28] (R = 0.796 on R-SIM)"
p27_text = paras[27].text
for run in paras[27].runs:
    txt = run.text
    # Fix R value and rename
    txt = (txt
        .replace('DeepMIF [27] (R = 0.784 on R-SIM)', 'DeepRSMA [28] (R = 0.796 on R-SIM)')
        .replace('DeepMIF [27] (R = 0.796 on R-SIM)', 'DeepRSMA [28] (R = 0.796 on R-SIM)')
        .replace('DeepMIF [27]', 'DeepRSMA [28]')
    )
    if txt != run.text:
        run.text = txt
        print("  Para 027: DeepMIF/R-value attribution fixed")

# Scan ALL body paragraphs for remaining [29] citations or EMMPTNet mentions
print("\nScanning for remaining EMMPTNet / [29] mentions...")
for i, p in enumerate(paras[:ref_start]):
    changed = False
    for run in p.runs:
        orig = run.text
        new = (orig
            .replace('EMMPTNet [29]', 'DeepMIF [27]')
            .replace('[29]', '')
            .replace('EMMPTNet', 'DeepMIF'))
        if new != orig:
            run.text = new
            changed = True
    if changed:
        print(f"  Para [{i}] updated: {paras[i].text[:100]}")

# ── 8. Save intermediate (before renumbering) ───────────────────────────────
TEMP = "Revisions/_temp_before_renumber.docx"
doc.save(TEMP)
print(f"\nIntermediate saved → {TEMP}")

# ── 9. Renumber citations ────────────────────────────────────────────────────
print("\nApplying citation renumbering...")
doc2 = Document(TEMP)
paras2 = list(doc2.paragraphs)

CITE_PAT = re.compile(r'\[(\d+(?:[,\s]*\d+)*(?:[–\-]\d+)?)\]')

def parse_cite_nums(s):
    s = s.strip()
    m = re.match(r'^(\d+)[–\-](\d+)$', s)
    if m:
        return list(range(int(m.group(1)), int(m.group(2)) + 1))
    return [int(x) for x in re.split(r'[,\s]+', s) if x.strip().isdigit()]

ref_start2 = next(i for i, p in enumerate(paras2)
                  if re.match(r'^[Rr]eferences?\s*$', p.text.strip()))

seen = {}; order = []
for i, p in enumerate(paras2[:ref_start2]):
    for m in CITE_PAT.finditer(p.text):
        for n in parse_cite_nums(m.group(1)):
            if n not in seen:
                seen[n] = i; order.append(n)

for tbl in doc2.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for m in CITE_PAT.finditer(p.text):
                    for n in parse_cite_nums(m.group(1)):
                        if n not in seen:
                            seen[n] = 'table'; order.append(n)

# Add uncited refs at end
all_ref_nums = []
for p in paras2[ref_start2 + 1:]:
    m = re.match(r'^\[(\d+)\]', p.text.strip())
    if m:
        all_ref_nums.append(int(m.group(1)))
uncited = [n for n in all_ref_nums if n not in seen]
if uncited:
    print(f"  Uncited refs (placed at end): {uncited}")
order.extend(uncited)

old_to_new = {old: new for new, old in enumerate(order, 1)}
print(f"  {len(order)} references total after EMMPTNet removal")

def placeholderize(text):
    def repl(m):
        inner = m.group(1)
        return '[' + re.sub(r'\d+', lambda nm: f'<{nm.group()}>', inner) + ']'
    return CITE_PAT.sub(repl, text)

def finalize(text):
    return re.sub(r'<(\d+)>', lambda m: str(old_to_new.get(int(m.group(1)), int(m.group(1)))), text)

for p in paras2[:ref_start2]:
    for run in p.runs:
        if '[' in run.text: run.text = placeholderize(run.text)
for p in paras2[:ref_start2]:
    for run in p.runs:
        if '<' in run.text: run.text = finalize(run.text)

for tbl in doc2.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if '[' in run.text: run.text = placeholderize(run.text)
for tbl in doc2.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if '<' in run.text: run.text = finalize(run.text)

# Reorder reference list
ref_entries = []
for p in paras2[ref_start2 + 1:]:
    txt = p.text.strip()
    m = re.match(r'^\[(\d+)\]', txt)
    if m:
        ref_entries.append((int(m.group(1)), p))

ref_entries_sorted = sorted(ref_entries, key=lambda x: old_to_new.get(x[0], 9999))
body2 = doc2.element.body
first_ref_pos2 = list(body2).index(ref_entries[0][1]._element)
for (_, p) in ref_entries:
    body2.remove(p._element)
for offset, (old_n, p) in enumerate(ref_entries_sorted):
    body2.insert(first_ref_pos2 + offset, p._element)

ns2 = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
for new_num, (old_n, p) in enumerate(ref_entries_sorted, 1):
    for t_elem in p._element.findall(f'.//{ns2}t'):
        if t_elem.text:
            new_text = re.sub(r'^\[' + str(old_n) + r'\]', f'[{new_num}]', t_elem.text)
            if new_text != t_elem.text:
                t_elem.text = new_text
                break

doc2.save(DOCX_OUT)
print(f"\nSaved → {DOCX_OUT}")

# ── Verify ───────────────────────────────────────────────────────────────────
print("\nVerification...")
doc3 = Document(DOCX_OUT)
paras3 = list(doc3.paragraphs)
ref_start3 = next(i for i, p in enumerate(paras3)
                  if re.match(r'^[Rr]eferences?\s*$', p.text.strip()))

seen3 = {}; max_n = 0; issues = 0
for i, p in enumerate(paras3[:ref_start3]):
    for m in CITE_PAT.finditer(p.text):
        for n in parse_cite_nums(m.group(1)):
            if n not in seen3:
                seen3[n] = i
                if n < max_n:
                    print(f"  ORDER ISSUE: [{n}] at para {i} after [{max_n}]"); issues += 1
                else:
                    max_n = n

total_refs = sum(1 for p in paras3[ref_start3+1:]
                 if re.match(r'^\[\d+\]', p.text.strip()))
print(f"  Citation order violations: {issues}")
print(f"  Total references: {total_refs}")
print(f"  Unique in-text citations: {len(seen3)}")

print("\nFinal reference list:")
for p in paras3[ref_start3+1:]:
    t = p.text.strip()
    if re.match(r'^\[\d+\]', t):
        print(f"  {t[:100]}")

print("\nTable 1:")
for ri, row in enumerate(doc3.tables[0].rows):
    print(f"  Row {ri}: {row.cells[0].text[:30]} | {row.cells[1].text[:15]} | {row.cells[4].text}")
