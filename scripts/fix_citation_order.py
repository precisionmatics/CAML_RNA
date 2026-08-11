#!/usr/bin/env python3
"""
fix_citation_order.py
Renumber all [N] citations so they appear in ascending first-appearance order.
Also reorders the reference list to match.
"""

import re
from docx import Document

DOCX_IN  = "Revisions/08_August_2026_Revised_CAML_RNA_manuscript.docx"
DOCX_OUT = "Revisions/09_August_2026_Revised_CAML_RNA_manuscript.docx"

doc = Document(DOCX_IN)
paras = list(doc.paragraphs)

# ── Citation pattern: [N], [N,M], [N-M], [N–M] ───────────────────────────────
CITE_PAT = re.compile(r'\[(\d+(?:[,\s]*\d+)*(?:[–\-]\d+)?)\]')

def parse_cite_nums(s):
    s = s.strip()
    m = re.match(r'^(\d+)[–\-](\d+)$', s)
    if m:
        return list(range(int(m.group(1)), int(m.group(2)) + 1))
    return [int(x) for x in re.split(r'[,\s]+', s) if x.strip().isdigit()]

# ── Step 1: find References section ──────────────────────────────────────────
ref_start = None
for i, p in enumerate(paras):
    if re.match(r'^[Rr]eferences?\s*$', p.text.strip()):
        ref_start = i
        break
print(f"References section: paragraph {ref_start} → '{paras[ref_start].text.strip()}'")

# ── Step 2: scan body paragraphs for first-appearance order ──────────────────
seen  = {}   # old_num → location
order = []   # old citation numbers in first-appearance order

for i, p in enumerate(paras[:ref_start]):
    for m in CITE_PAT.finditer(p.text):
        for n in parse_cite_nums(m.group(1)):
            if n not in seen:
                seen[n] = ('body', i)
                order.append(n)

# Also scan table cells
for ti, tbl in enumerate(doc.tables):
    for ri, row in enumerate(tbl.rows):
        for ci, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                for m in CITE_PAT.finditer(p.text):
                    for n in parse_cite_nums(m.group(1)):
                        if n not in seen:
                            seen[n] = ('table', ti)
                            order.append(n)

# Refs in reference list but not in body/tables → append at end
all_ref_nums = []
for p in paras[ref_start + 1:]:
    m = re.match(r'^\[(\d+)\]', p.text.strip())
    if m:
        all_ref_nums.append(int(m.group(1)))

uncited = [n for n in all_ref_nums if n not in seen]
print(f"\nUncited refs (not in body/tables): {uncited}")
order.extend(uncited)

# ── Step 3: build old→new map ─────────────────────────────────────────────────
old_to_new = {old_n: new_idx for new_idx, old_n in enumerate(order, 1)}

print(f"\nRenumbering map ({len(order)} total):")
for new_idx, old_n in enumerate(order, 1):
    print(f"  [{old_n:2d}] → [{new_idx:2d}]  ({seen.get(old_n, 'ref-list-only')})")

# ── Step 4: apply renumbering via placeholder trick ───────────────────────────
# Pass 1: replace digit sequences inside [...] with ‹N› placeholders
# This preserves format: [14-16] → [‹14›-‹16›], [3,4] → [‹3›,‹4›]

def placeholderize(text):
    def repl(m):
        inner = m.group(1)
        new_inner = re.sub(r'\d+', lambda nm: f'‹{nm.group()}›', inner)
        return f'[{new_inner}]'
    return CITE_PAT.sub(repl, text)

def finalize(text):
    def repl(m):
        old_n = int(m.group(1))
        return str(old_to_new.get(old_n, old_n))
    return re.sub(r'‹(\d+)›', repl, text)

# Apply to body paragraphs
for p in paras[:ref_start]:
    for run in p.runs:
        if '[' in run.text:
            run.text = placeholderize(run.text)
for p in paras[:ref_start]:
    for run in p.runs:
        if '‹' in run.text:
            run.text = finalize(run.text)

# Apply to table cells
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if '[' in run.text:
                        run.text = placeholderize(run.text)
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if '‹' in run.text:
                        run.text = finalize(run.text)

# ── Step 5: reorder reference list ───────────────────────────────────────────
ref_entries = []  # (old_num, para_obj)
for p in paras[ref_start + 1:]:
    txt = p.text.strip()
    m = re.match(r'^\[(\d+)\]', txt)
    if m:
        ref_entries.append((int(m.group(1)), p))

print(f"\nFound {len(ref_entries)} reference list entries")

# Sort by new numbering
ref_entries_sorted = sorted(ref_entries, key=lambda x: old_to_new.get(x[0], 9999))

# Reorder XML elements in body
body = doc.element.body
first_ref_pos = list(body).index(ref_entries[0][1]._element)

for (_, p) in ref_entries:
    body.remove(p._element)

for offset, (old_n, p) in enumerate(ref_entries_sorted):
    body.insert(first_ref_pos + offset, p._element)

# Renumber [old_N] prefix of each reference entry to [new sequential number]
ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
for new_num, (old_n, p) in enumerate(ref_entries_sorted, 1):
    for t_elem in p._element.findall(f'.//{ns}t'):
        if t_elem.text:
            new_text = re.sub(r'^\[' + str(old_n) + r'\]', f'[{new_num}]', t_elem.text)
            if new_text != t_elem.text:
                t_elem.text = new_text
                break

# ── Save ──────────────────────────────────────────────────────────────────────
doc.save(DOCX_OUT)
print(f"\nSaved → {DOCX_OUT}")

# ── Verify ────────────────────────────────────────────────────────────────────
print("\nVerification — checking citation order in saved file...")
doc2 = Document(DOCX_OUT)
paras2 = list(doc2.paragraphs)
ref_start2 = None
for i, p in enumerate(paras2):
    if re.match(r'^[Rr]eferences?\s*$', p.text.strip()):
        ref_start2 = i
        break

seen2 = {}
max_so_far = 0
issues = 0
for i, p in enumerate(paras2[:ref_start2]):
    for m in CITE_PAT.finditer(p.text):
        for n in parse_cite_nums(m.group(1)):
            if n not in seen2:
                seen2[n] = i
                if n < max_so_far:
                    print(f"  OUT OF ORDER: [{n}] at para {i} after [{max_so_far}]")
                    issues += 1
                else:
                    max_so_far = n

# Also check tables
for tbl in doc2.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for m in CITE_PAT.finditer(p.text):
                    for n in parse_cite_nums(m.group(1)):
                        if n not in seen2:
                            seen2[n] = 'table'
                            if n < max_so_far:
                                print(f"  OUT OF ORDER (table): [{n}] after [{max_so_far}]")
                                issues += 1
                            else:
                                max_so_far = n

# Check reference list is in order
print("\nReference list order:")
for p in paras2[ref_start2 + 1:]:
    txt = p.text.strip()
    m = re.match(r'^\[(\d+)\]', txt)
    if m:
        print(f"  {txt[:90]}")

if issues == 0:
    print(f"\nAll {len(seen2)} citations in correct ascending order!")
else:
    print(f"\n{issues} order violations remain — check output")
